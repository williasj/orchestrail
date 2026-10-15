# -*- coding: utf-8 -*-
"""Independent grader for the sales-report benchmark (CLI: -i CSV -o XLSX, --generate-sample).
Usage: python3 grade_sales.py sales_report.py"""
import os, random, subprocess, sys, tempfile
import openpyxl, pandas as pd
script = os.path.abspath(sys.argv[1])
score = []
def chk(n, ok):
    score.append(bool(ok)); print('PASS' if ok else 'FAIL', n)
def run(*args, cwd):
    return subprocess.run([sys.executable, script, *args], cwd=cwd, capture_output=True, text=True, timeout=120)
d = tempfile.mkdtemp()
r = run('--generate-sample', '-i', 'sample.csv', cwd=d)
chk('sample generator works', r.returncode == 0 and os.path.exists(os.path.join(d, 'sample.csv')))
# our own data with known answers
rng = random.Random(7); rows = []
for i in range(300):
    rows.append({'date': f'2026-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}', 'region': rng.choice(['North', 'South', 'East', 'West']),
                 'product': rng.choice(['A', 'B', 'C', 'D', 'E']), 'units': rng.randint(1, 20), 'unit_price': round(rng.uniform(1, 50), 2)})
df = pd.DataFrame(rows); df.to_csv(os.path.join(d, 'mine.csv'), index=False)
df['rev'] = df.units * df.unit_price
r = run('-i', 'mine.csv', '-o', 'out.xlsx', cwd=d)
chk('report runs on real data', r.returncode == 0)
wb = openpyxl.load_workbook(os.path.join(d, 'out.xlsx'))
names = [s.lower() for s in wb.sheetnames]
chk('three sheets By Region / By Month / Top Products', all(any(k in n for n in names) for k in ('region', 'month', 'top')))
def sheet(key): return wb[[s for s in wb.sheetnames if key in s.lower()][0]]
def nums(ws): return [c for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, (int, float))]
reg = df.groupby('region').rev.sum()
chk('region revenue correct', all(any(abs(v - x) < 0.01 for x in nums(sheet('region'))) for v in reg.values))
mon = df.assign(m=df.date.str[:7]).groupby('m').rev.sum()
chk('monthly revenue correct', all(any(abs(v - x) < 0.01 for x in nums(sheet('month'))) for v in mon.values))
top = df.groupby('product').rev.sum().sort_values(ascending=False)
ws = sheet('top'); txt = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None]
chk('top 3 products correct and only 3', [p for p in txt if p in top.index][:3] == list(top.index[:3]) and sum(1 for p in txt if p in top.index) == 3)
chk('bar chart on By Region', any(type(ch).__name__ == 'BarChart' for ch in getattr(sheet('region'), '_charts', [])))
r = run('-i', 'missing.csv', '-o', 'x.xlsx', cwd=d)
chk('missing file: exit 1 with message', r.returncode == 1 and (r.stderr.strip() or r.stdout.strip()) and 'Traceback' not in r.stderr)
open(os.path.join(d, 'empty.csv'), 'w').close()
r = run('-i', 'empty.csv', '-o', 'x.xlsx', cwd=d)
chk('empty file: exit 1 with message', r.returncode == 1 and 'Traceback' not in r.stderr)
open(os.path.join(d, 'hdr.csv'), 'w').write('date,region,product,units,unit_price\n')
r = run('-i', 'hdr.csv', '-o', 'x.xlsx', cwd=d)
chk('header-only file: exit 1, no traceback', r.returncode == 1 and 'Traceback' not in r.stderr)
print(f'SCORE {sum(score)}/{len(score)}')
